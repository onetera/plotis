# :coding: utf-8

import re
import yaml
import openai

from langchain.chat_models import ChatOpenAI
from langchain_core.prompts import SystemMessagePromptTemplate , HumanMessagePromptTemplate
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser
from langchain_core import output_parsers 
from langchain_community.document_loaders import PyPDFLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.runnables import RunnablePassthrough
from langchain.schema import Document

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment

import requests
from io import BytesIO
import shortuuid

import db_conn



class PreprodAI:
    def __init__(self ):
        with open( './config.yml' ) as f:
            data = yaml.load( f, Loader = yaml.FullLoader )
            self.api_key = data['api_key']
        
        self.oai_client = openai.OpenAI( 
                                    api_key = self.api_key ,
                                    #**params,
        )

        self.sys_temp = SystemMessagePromptTemplate.from_template(
                    ' 이 시스템은 한국 영화 시나리오 작가이다.'
        )
        self.db = db_conn.DBconn()

        self.synop = ''
        self.loc = None
        self.scene_list = []
        self.scenario = ''

        self.scene_numbers = []
    
    def client(self, temperature):
        return ChatOpenAI(
                    model = 'gpt-4o',
                    # model = 'gpt-4o-mini',
                    api_key = self.api_key,
                    temperature=temperature
        )
        
    def combine_scene( self ):
        loc = self.create_location( self.synop )
        self.loc = loc.choices[0].message.content


    def create_synop( self, *key ):    
        key_list = key
        key_join = ','.join( key_list  )
        search_msg = '{key_join} 이런 내용의 시놉시스를 작성해줘.'
        search_msg += '제목같은 다른 내용은 출력하지 말고 시놉시스의 내용만 출력'
        chain = self.chain( search_msg )
        response = chain.invoke({'key_join':key_join}  )
        self.synop = response
        self.db.insert_synop( response, key_join )
        return response

    def create_location( self, min = 80, max = 120, synop = '' ):
        if synop:
            synop = synop
        else:
            if self.synop:
                synop = self.synop
            else:
                return

        search_msg = '{synop}'
        search_msg += '이 시스템은 한국 영화 시나리오 작가이다'
        search_msg += '이 시놉시스를 이용 해서 기승전결있는  {min}~{max}개의 장면을 만들어줘'
        search_msg += '장면번호는 숫자만 , 장소, 간단한 설명의 순서로 컴마로 구분된  텍스트 형태 데이이터로 만들어줘'
        search_msg += '데이터 형식 이름 같이  다른건 아무것도 출력하지 말고 오직 생성된 중첩 데이터만 출력해줘.'
        search_msg += '장면이 {min}개 이하면 좀더 계산해서 기승전결이 있는 {min}개~{max}개의 장면이 되도록 작성해줘.'

        chain = self.chain( search_msg )
        response = chain.invoke( 
                        {'synop':synop, 'min' : min, 'max':max } 
        )
        loc_list = []
        for row in response.split('\n'):
            loc_list.append( row.split(',') )
        self.loc = loc_list
        return response

    def write_scene( self ):
        ## location 기반으로 작성
        synop_idx = self.db.search_synop_idx(self.synop)

        for loc in self.loc:
            print( 'loc : ', loc )
            search_msg =  '이 시스템은 한국 영화 시나리오 작가이다.'
            search_msg += '이 장면의 번호는 {num}이다.'
            search_msg += '이 장면의 장소는 {location}이다.'
            search_msg += '이 장면의 내용은 {desc}이다.'
            search_msg += '이 정보를 이용해서 상세한 시나리오를 작성해줘.'
            search_msg += '장면3의 예시) ### 장면 3: 장면의 장소'

            chain = self.chain( search_msg )
            response = chain.invoke( 
                        {   
                            'num'     : loc[0],
                            'location': loc[1],
                            'desc'    : loc[2],
                        }

            )
            self.scene_list.append( [ loc[1], response] )
            self.scenario += response
            self.scenario += '\n'
        
        self.db.insert_scenario(self.scenario, synop_idx)

    def write_conti( self, scenario, scenario_idx ):
        scene_pattern = re.compile(r'\#\#\# 장면 \d+:')

        scenes = scene_pattern.split(scenario)
        scene_titles = scene_pattern.findall(scenario)

        scene_list = [f"{title.strip()} {scene.strip()}" for title, scene in zip(scene_titles, scenes[1:])]

        prompt_ori = '이 시스템은 영화 콘티작가 입니다.다음 시나리오 일부를 실사 이미지화 해주세요.'
        prompt_ori += "다음에 오는 내용에서 중요 장면 1개를 선정한 후 반드시 아래의 스타일에 맞춰 그려주세요"
        prompt_ori += "만약 컨텐츠 정책이 위반되었다면 반드시 정책을 준수해서 다시 장면 선정을 한 후 이미지를 재생성해주세요."
        prompt_ori += "스타일 : 흑색의 크로키 스타일 스케치, 본질과 감정을 담아낸 부드럽고 표현적인 선, 간결한 배경, 자연스러운 느낌"

        load_conti = self.db.load_conti( scenario_idx )

        for data in scene_list:
            prompt = prompt_ori

            prompt += f"내용 : {data}"
            prompt += "만약 컨텐츠 정책이 위반되었다면 반드시 정책을 준수할 때까지 다시 장면을 선정해서 이미지를 재생성해주세요."

            response = self.oai_client.images.generate(
                model = 'dall-e-3',
                prompt = prompt,
                n = 1,
                size = '1024x1024',
                quality = 'standard',
                style = 'natural'
            )
            image_url = response.data[0].url
            image_url_get = requests.get(image_url)

            img_uid = str(shortuuid.uuid())
            img_path = f'./tmp/{img_uid}.png'
            with open(img_path, 'wb') as f:
                f.write(image_url_get.content)
            
            if not load_conti:
                self.db.insert_conti( data, img_path, scenario_idx )
            else:
                self.db.update_conti( data, img_path, scenario_idx )
            
    def save_conti( self, scenario_idx ):
        wb = Workbook()
        ws = wb.active

        contis = self.db.load_conti( scenario_idx )

        for row, conti_data in enumerate( contis ):
            scene = conti_data[1]
            img_path = conti_data[2]

            cell = ws.cell( row = row+1, column = 1 , value = scene )
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

            with open(img_path, 'rb') as img_file:
                img = Image(BytesIO(img_file.read()))
            img.height = 320
            img.width = 320
            ws.add_image( img, 'B{}'.format( row + 1 ) )
            ws.row_dimensions[ row+1 ].height = 320
            ws.column_dimensions[ 'A' ].width = 90

        conti_path = './tmp/conti.xlsx'
        wb.save(  conti_path )
        return conti_path
           

        

    def chain( self, search_msg , parser = StrOutputParser() ):
        human_temp = HumanMessagePromptTemplate.from_template( search_msg )
        chat_prompt = ChatPromptTemplate.from_messages(
            [
                self.sys_temp,
                human_temp,
            ]
        )
        chain = chat_prompt|self.client(0.5)| output_parsers.StrOutputParser()
        return chain

        

    def remove_page_numbers( self, text ):
        text = re.sub(r'\s*-\s*\d+\s*-\s*', '', text, flags=re.MULTILINE)
        return text

    def combine_pages_with_overlap( self, pages, pages_per_combination=10, overlap=150 ):
        combined_pages = []
        num_pages = len(pages)
        
        for i in range(0, num_pages, pages_per_combination):
            combined_text = ''
            
            if i > 0:
                previous_overlap = pages[i-1][-overlap:]
                combined_text += previous_overlap + ' '
            
            combined_text += ' '.join(pages[i:i + pages_per_combination])
            
            if i + pages_per_combination < num_pages:
                next_overlap = pages[i + pages_per_combination][:overlap]
                combined_text += ' ' + next_overlap
            
            combined_pages.append(combined_text)
        
        return combined_pages

    def add_blank_lines_around_numbers( self, text ):
        pattern = r'(?<!\n)\n(?:\d{0,3}|\#|\s*SCENE\s*|\s*씬\s*|\s*S#\s*|\s*장면\s*)\d+.*?(?=\n|\Z)'
        replacement = r'\n\g<0>\n'
        updated_text = re.sub(pattern, replacement, text, flags=re.DOTALL | re.IGNORECASE)
        return updated_text

    def format_documents( self, docs ):
        if '\n' not in docs:
            for doc in docs:
                combined_text = doc.page_content.replace('. ', '.\n')
                combined_text = self.add_blank_lines_around_numbers(combined_text)
        else:
            combined_text = '\n'.join(doc.page_content for doc in docs)
        return combined_text

    def parse_scene_results( self, results ):
        parsed_results = []
        pattern = re.compile(r'(\d{1,3}),\s*(.+),\s*(.+)', re.DOTALL | re.IGNORECASE)
        for res in results:
            lines = res.split('\n')
            for line in lines:
                line = line.strip()
                if line:
                    line = line.strip('[]')
                    match = pattern.match(line)
                    if match:
                        scene_number = match.group(1).strip()
                        location = match.group(2).strip()
                        summary = match.group(3).strip()
                        if scene_number not in self.scene_numbers:
                            parsed_results.append([scene_number, location, summary])
                            self.scene_numbers.append(scene_number)
        return parsed_results

    def find_location_from_pdf( self, pdf_filepath ):
        loader = PyPDFLoader(pdf_filepath)
        pages = loader.load()

        cleaned_pages = [self.remove_page_numbers(page.page_content) for page in pages]

        combined_pages = self.combine_pages_with_overlap(cleaned_pages, pages_per_combination=5, overlap=150)

        cleaned_documents = [
            Document(page_content=content, metadata=pages[i].metadata)
            for i, content in enumerate(combined_pages)
        ]

        text_splitter = RecursiveCharacterTextSplitter(chunk_size=50000, chunk_overlap=200)
        splits = text_splitter.split_documents(cleaned_documents)

        template = "당신은 시나리오 전문가입니다."
        template += "다음 시나리오를 바탕으로 각 씬을 다음 형식으로 요약하세요 : [씬 넘버, 메인 장소, 씬 내용 1줄요약]."
        template += "씬 내용 1줄요약은 '~하는 장면'으로 요약해 주세요. 추가 설명 없이 형식에 맞춰서만 작성해 주세요."
        template += "응답은 반드시 한국어로 작성하세요."
        template += " **씬 넘버는 시나리오 내용에 포함된 번호를 반드시 그대로 사용해 주세요.**"
        template += "아래는 예시입니다:"
        template += "예시 시나리오:"
        template += "씬 46: 공원"
        template += "주인공이 친구와 만나 이야기를 나눈다."
        template += "씬 47: 카페"
        template += "주인공이 커피를 마시며 서류를 정리한다."
        template += "응답 예시:"
        template += "[46, 공원, 주인공이 친구와 만나 이야기를 나누는 장면]"
        template += "[47, 카페, 주인공이 커피를 마시며 서류를 정리하는 장면]"
        template += "내용: {context}"

        prompt = ChatPromptTemplate.from_template(template)

        rag_chain = (
            {'context': RunnablePassthrough()}
            | prompt
            | self.client(0)
            | StrOutputParser()
        )

        final_results = []
        for i in range(len(splits)):
            results = []
            chunk_docs = [splits[i]]
            formatted_documents = self.format_documents(chunk_docs)
            result = rag_chain.invoke({"context": formatted_documents})
            results.append(result.strip())
            parsed_results = self.parse_scene_results(results)
            final_results.extend(parsed_results)

        return final_results

